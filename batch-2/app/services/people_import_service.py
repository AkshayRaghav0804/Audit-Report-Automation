"""Preview-first Excel and CSV imports for the application people directory."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from collections import Counter
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models.enums import ImportKind, ImportStatus
from app.database.models.import_data import (
    MasterDataImportBatch,
    MasterDataImportError,
)
from app.database.session import DatabaseManager, database
from app.repositories.directory_repository import DirectoryRepository
from app.schemas.directory import (
    DirectoryPerson,
    ImportIssue,
    ImportRowAction,
    PeopleImportPreview,
    PeopleImportRow,
    PersonUpsert,
)
from app.services.authorization_service import require_role_by_id
from app.services.directory_service import (
    ROLE_CATALOG,
    directory_person_for_user,
    ensure_role_catalog,
    upsert_person,
)

REQUIRED_HEADERS = frozenset({"employee_code", "full_name", "email", "roles"})
MAX_IMPORT_BYTES = 5_000_000
MAX_IMPORT_ROWS = 10_000
MAX_IMPORT_COLUMNS = 64
MAX_XLSX_ARCHIVE_MEMBERS = 1_000
MAX_XLSX_UNCOMPRESSED_BYTES = 50_000_000
MAX_XLSX_MEMBER_BYTES = 20_000_000
HEADER_ALIASES = {
    "name": "full_name",
    "display_name": "full_name",
    "employee_id": "employee_code",
    "icai": "icai_membership_number",
    "membership_number": "icai_membership_number",
}


class UnsupportedImportFileError(ValueError):
    """Raised when a people import file is not an accepted workbook or CSV."""


class PeopleImportService:
    def __init__(self, manager: DatabaseManager = database) -> None:
        self.manager = manager

    async def preview(self, filename: str, content: bytes) -> PeopleImportPreview:
        normalized_filename = _normalize_filename(filename)
        records = _read_records(normalized_filename, content)
        async with self.manager.session() as session:
            return await _build_preview(session, normalized_filename, content, records)

    async def commit(self, preview: PeopleImportPreview, *, uploaded_by_user_id: UUID) -> UUID:
        if not preview.can_commit:
            raise ValueError("Resolve all blocking import errors before committing.")
        _normalize_filename(preview.filename)
        if len(preview.rows) > MAX_IMPORT_ROWS:
            raise ValueError(f"People imports cannot contain more than {MAX_IMPORT_ROWS:,} rows.")
        async with self.manager.session() as session:
            async with session.begin():
                try:
                    await require_role_by_id(session, uploaded_by_user_id, "administrator")
                except PermissionError as error:
                    raise PermissionError(
                        "Only an active application administrator can import people."
                    ) from error
                await ensure_role_catalog(session)
                batch = MasterDataImportBatch(
                    kind=ImportKind.PEOPLE,
                    filename=preview.filename,
                    file_sha256=preview.file_sha256,
                    status=ImportStatus.PREVIEWED,
                    uploaded_by_user_id=uploaded_by_user_id,
                    preview_json={
                        "rows": len(preview.rows),
                        "inserted": preview.inserted_count,
                        "updated": preview.updated_count,
                        "unchanged": preview.unchanged_count,
                    },
                    inserted_count=preview.inserted_count,
                    updated_count=preview.updated_count,
                    unchanged_count=preview.unchanged_count,
                    warning_count=preview.warning_count,
                    error_count=preview.error_count,
                )
                session.add(batch)
                await session.flush()
                for row in preview.rows:
                    if row.person is not None and row.action is not ImportRowAction.ERROR:
                        await upsert_person(session, row.person)
                    for issue in row.issues:
                        session.add(
                            MasterDataImportError(
                                batch_id=batch.id,
                                sheet_name="People",
                                row_number=row.row_number,
                                field_name=issue.field_name,
                                error_code=issue.code,
                                message=issue.message,
                                raw_row_json={},
                            )
                        )
                batch.status = ImportStatus.COMMITTED
                from app.database.models.mixins import utc_now

                batch.committed_at = utc_now()
                return batch.id


async def _build_preview(
    session: AsyncSession,
    filename: str,
    content: bytes,
    records: tuple[tuple[int, Mapping[str, Any]], ...],
) -> PeopleImportPreview:
    repository = DirectoryRepository(session)
    await ensure_role_catalog(session)
    parsed: list[tuple[int, PersonUpsert | None, list[ImportIssue]]] = []
    for row_number, record in records:
        issues: list[ImportIssue] = []
        try:
            person = _person_from_record(record)
        except ValidationError as error:
            person = None
            issues.extend(
                ImportIssue(
                    row_number=row_number,
                    field_name=str(problem["loc"][-1]) if problem["loc"] else None,
                    code="invalid_value",
                    message=str(problem["msg"]),
                )
                for problem in error.errors()
            )
        except ValueError as error:
            person = None
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    code="invalid_row",
                    message=str(error),
                )
            )
        parsed.append((row_number, person, issues))

    people = tuple(person for _, person, _ in parsed if person is not None)
    code_counts = Counter(person.employee_code for person in people)
    email_counts = Counter(person.email for person in people)
    credential_counts = Counter(
        person.icai_membership_number for person in people if person.icai_membership_number
    )
    name_counts = Counter(person.display_name.casefold() for person in people)

    rows: list[PeopleImportRow] = []
    for row_number, person, issues in parsed:
        if person is None:
            rows.append(
                PeopleImportRow(
                    row_number=row_number,
                    person=None,
                    action=ImportRowAction.ERROR,
                    issues=tuple(issues),
                )
            )
            continue
        if code_counts[person.employee_code] > 1:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    field_name="employee_code",
                    code="duplicate_employee_code",
                    message="Employee code is repeated in the import file.",
                )
            )
        if email_counts[person.email] > 1:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    field_name="email",
                    code="duplicate_email",
                    message="Email is repeated in the import file.",
                )
            )
        if person.icai_membership_number and credential_counts[person.icai_membership_number] > 1:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    field_name="icai_membership_number",
                    code="duplicate_icai",
                    message="ICAI membership number is repeated in the import file.",
                )
            )
        if "partner" in person.roles and not person.icai_membership_number:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    field_name="icai_membership_number",
                    code="partner_icai_required",
                    message="Partners must have an ICAI membership number.",
                )
            )
        unknown_roles = set(person.roles).difference(ROLE_CATALOG)
        if unknown_roles:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    field_name="roles",
                    code="unknown_role",
                    message="Unknown roles: " + ", ".join(sorted(unknown_roles)),
                )
            )
        if name_counts[person.display_name.casefold()] > 1:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    field_name="full_name",
                    code="duplicate_display_name",
                    message=(
                        "Another imported person has the same name. Both are allowed and "
                        "will be distinguished by email and employee code."
                    ),
                    blocking=False,
                )
            )

        by_code = await repository.find_user_by_employee_code(person.employee_code)
        by_email = await repository.find_user_by_email(person.email)
        if by_code and by_email and by_code.id != by_email.id:
            issues.append(
                ImportIssue(
                    row_number=row_number,
                    code="identity_conflict",
                    message="Employee code and email match different existing users.",
                )
            )
        if person.icai_membership_number:
            credential = await repository.find_credential("ICAI", person.icai_membership_number)
            matched_user = by_code or by_email
            if credential and (matched_user is None or credential.user_id != matched_user.id):
                issues.append(
                    ImportIssue(
                        row_number=row_number,
                        field_name="icai_membership_number",
                        code="icai_conflict",
                        message="ICAI membership number belongs to another directory user.",
                    )
                )
        blocking = any(issue.blocking for issue in issues)
        if blocking:
            action = ImportRowAction.ERROR
        elif by_code or by_email:
            matched_user = by_code or by_email
            if matched_user is None:  # pragma: no cover - narrowed by the branch above
                raise RuntimeError("Unable to resolve the matched directory user.")
            existing = await directory_person_for_user(session, matched_user.id)
            action = (
                ImportRowAction.UNCHANGED
                if existing is not None and _person_matches(person, existing)
                else ImportRowAction.UPDATE
            )
        else:
            action = ImportRowAction.INSERT
        rows.append(
            PeopleImportRow(
                row_number=row_number,
                person=person,
                action=action,
                issues=tuple(issues),
            )
        )
    return PeopleImportPreview(
        filename=filename,
        file_sha256=hashlib.sha256(content).hexdigest(),
        rows=tuple(rows),
    )


def _read_records(filename: str, content: bytes) -> tuple[tuple[int, Mapping[str, Any]], ...]:
    if not content:
        raise UnsupportedImportFileError("The people import file is empty.")
    if len(content) > MAX_IMPORT_BYTES:
        raise UnsupportedImportFileError(
            f"People import files must be {MAX_IMPORT_BYTES // 1_000_000} MB or smaller."
        )
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            raw_headers = tuple(reader.fieldnames or ())
            normalized_headers = tuple(_normalize_header(header) for header in raw_headers)
            _validate_headers(normalized_headers)
            csv_headers = dict(zip(raw_headers, normalized_headers, strict=True))
            records: list[tuple[int, Mapping[str, Any]]] = []
            for row_number, row in enumerate(reader, start=2):
                if not any(value and str(value).strip() for value in row.values()):
                    continue
                if len(records) >= MAX_IMPORT_ROWS:
                    raise ValueError(
                        f"People imports cannot contain more than {MAX_IMPORT_ROWS:,} rows."
                    )
                records.append((row_number, _normalize_record(row, csv_headers)))
            return tuple(records)
        except (UnicodeError, csv.Error) as error:
            raise UnsupportedImportFileError(
                "The CSV file could not be read as UTF-8 data."
            ) from error
    if suffix == ".xlsx":
        _validate_xlsx_archive(content)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook["People"] if "People" in workbook.sheetnames else workbook.active
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, ())
            xlsx_headers = tuple(_normalize_header(str(value or "")) for value in header_row)
            _validate_headers(xlsx_headers)
            records = []
            for row_number, values in enumerate(rows, start=2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                if len(records) >= MAX_IMPORT_ROWS:
                    raise ValueError(
                        f"People imports cannot contain more than {MAX_IMPORT_ROWS:,} rows."
                    )
                records.append(
                    (
                        row_number,
                        {
                            normalized: values[index] if index < len(values) else None
                            for index, normalized in enumerate(xlsx_headers)
                        },
                    )
                )
            return tuple(records)
        finally:
            workbook.close()
    raise UnsupportedImportFileError("Upload an .xlsx or .csv people file.")


def _normalize_filename(filename: str) -> str:
    normalized = re.split(r"[\\/]", str(filename))[-1].strip()
    if not normalized or any(ord(character) < 32 for character in normalized):
        raise UnsupportedImportFileError("The people import file must have a valid name.")
    if len(normalized) > 255:
        raise UnsupportedImportFileError("The people import filename is too long.")
    return normalized


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
    except (OSError, zipfile.BadZipFile) as error:
        raise UnsupportedImportFileError("The XLSX file is not a valid Excel workbook.") from error
    if len(members) > MAX_XLSX_ARCHIVE_MEMBERS:
        raise UnsupportedImportFileError("The XLSX file contains too many package entries.")
    if any(member.flag_bits & 0x1 for member in members):
        raise UnsupportedImportFileError("Encrypted XLSX files are not supported.")
    if any(member.file_size > MAX_XLSX_MEMBER_BYTES for member in members):
        raise UnsupportedImportFileError("The XLSX file expands beyond the safe import limit.")
    if sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise UnsupportedImportFileError("The XLSX file expands beyond the safe import limit.")


def _normalize_header(header: str) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", str(header).strip().casefold()).strip("_")
    return HEADER_ALIASES.get(key, key)


def _validate_headers(headers: Iterable[str]) -> None:
    normalized = tuple(headers)
    if len(normalized) > MAX_IMPORT_COLUMNS:
        raise ValueError(f"People imports cannot contain more than {MAX_IMPORT_COLUMNS} columns.")
    if not normalized or any(not header for header in normalized):
        raise ValueError("Every people import column must have a header.")
    duplicates = sorted(header for header, count in Counter(normalized).items() if count > 1)
    if duplicates:
        raise ValueError("Duplicate columns: " + ", ".join(duplicates))
    available = set(normalized)
    missing = REQUIRED_HEADERS.difference(available)
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(sorted(missing)))


def _normalize_record(record: Mapping[str, Any], headers: Mapping[str, str]) -> dict[str, Any]:
    return {headers.get(str(header), str(header)): value for header, value in record.items()}


def _person_from_record(record: Mapping[str, Any]) -> PersonUpsert:
    roles_text = str(record.get("roles") or "")
    roles = tuple(role.strip() for role in re.split(r"[;,]", roles_text) if role.strip())
    return PersonUpsert(
        employee_code=str(record.get("employee_code") or ""),
        display_name=str(record.get("full_name") or ""),
        email=str(record.get("email") or ""),
        department=_optional_text(record.get("department")),
        roles=roles,
        icai_membership_number=_optional_text(record.get("icai_membership_number")),
        active=_parse_active(record.get("active")),
    )


def _optional_text(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def _parse_active(value: Any) -> bool:
    if value is None or str(value).strip() == "":
        return True
    normalized = str(value).strip().casefold()
    if normalized in {"1", "true", "yes", "active"}:
        return True
    if normalized in {"0", "false", "no", "inactive"}:
        return False
    raise ValueError("Active must be yes/no, true/false, 1/0, or active/inactive.")


def _person_matches(person: PersonUpsert, existing: DirectoryPerson) -> bool:
    return (
        existing.employee_code == person.employee_code
        and existing.display_name == person.display_name
        and existing.email == person.email
        and existing.department == person.department
        and set(existing.roles) == set(person.roles)
        and existing.icai_membership_number == person.icai_membership_number
        and existing.active is person.active
    )
