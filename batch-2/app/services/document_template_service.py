"""Load report definitions from a catalog of standalone Word templates."""

from __future__ import annotations

import hashlib
import json
import logging
import re
import time
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, cast

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models.report_template import (
    PartnerReference,
    ReportContentSegment,
    ReportNumbering,
    ReportNumberingLevel,
    ReportParagraph,
    ReportSection,
    ReportTable,
    ReportTableCell,
    ReportTemplate,
)
from app.services.source_errors import SourceDataError

LOGGER = logging.getLogger(__name__)
EXPECTED_REPORT_COUNT = 7
REPORT_START = "independent auditor's report"
SOURCE_READ_ATTEMPTS = 3
SOURCE_READ_RETRY_SECONDS = 0.2
CATALOG_SCHEMA_VERSION = 1
EDITABLE_BLUE_RGB = "00B0F0"
CONDITIONAL_GREEN_RGB = frozenset({"4EA72E", "92D050"})


@dataclass(frozen=True, slots=True)
class _FileSignature:
    path: str
    modified_ns: int
    size: int


@dataclass(frozen=True, slots=True)
class _CatalogTemplate:
    id: str
    name: str
    source_order: int
    path: Path
    sha256: str


@dataclass(frozen=True, slots=True)
class _TemplateCatalog:
    source_file: str
    templates: tuple[_CatalogTemplate, ...]
    partners: tuple[PartnerReference, ...]


class DocumentTemplateService:
    """Parse individual Word templates and invalidate cache when the catalog changes."""

    def __init__(self, source_path: Path) -> None:
        self.source_path = source_path
        self._cached_signature: tuple[_FileSignature, ...] | None = None
        self._cached_templates: tuple[ReportTemplate, ...] = ()
        self._cached_partners: tuple[PartnerReference, ...] = ()

    def load_templates(self) -> tuple[ReportTemplate, ...]:
        catalog = self._read_catalog()
        signature = self._signature(catalog)
        if signature == self._cached_signature:
            return self._cached_templates

        templates: list[ReportTemplate] = []
        for entry in catalog.templates:
            document = self._read_document(entry.path, entry.sha256)
            report_blocks = self._split_report_blocks(document)
            if len(report_blocks) != 1:
                raise SourceDataError(
                    f"Standalone template '{entry.path.name}' must contain exactly one "
                    f"report body; found {len(report_blocks)}."
                )
            sections = self._extract_sections(entry.id, report_blocks[0])
            if not sections:
                raise SourceDataError(
                    f"Standalone template '{entry.path.name}' did not contain usable sections."
                )
            templates.append(
                ReportTemplate(
                    id=entry.id,
                    name=entry.name,
                    source_file=entry.path.name,
                    source_order=entry.source_order,
                    sections=sections,
                )
            )

        self._cached_signature = signature
        self._cached_templates = tuple(templates)
        self._cached_partners = catalog.partners
        LOGGER.info(
            "Loaded %d standalone report templates, %d sections, and %d partner "
            "references from %s (generated from %s)",
            len(templates),
            sum(len(template.sections) for template in templates),
            len(catalog.partners),
            self.source_path.name,
            catalog.source_file,
        )
        return self._cached_templates

    def load_partners(self) -> tuple[PartnerReference, ...]:
        """Return standardized partner names and ICAI numbers from the catalog."""

        self.load_templates()
        return self._cached_partners

    def _read_catalog(self) -> _TemplateCatalog:
        try:
            payload = self.source_path.read_text(encoding="utf-8-sig")
        except FileNotFoundError as error:
            raise SourceDataError(
                f"Report template catalog is missing: {self.source_path}"
            ) from error
        except OSError as error:
            raise SourceDataError(
                f"Unable to read report template catalog '{self.source_path.name}': {error}"
            ) from error

        try:
            raw: object = json.loads(payload)
        except json.JSONDecodeError as error:
            raise SourceDataError(
                f"Report template catalog '{self.source_path.name}' is not valid JSON: {error}"
            ) from error

        root = _mapping(raw, "catalog")
        if root.get("schema_version") != CATALOG_SCHEMA_VERSION:
            raise SourceDataError(
                f"Report template catalog must use schema version {CATALOG_SCHEMA_VERSION}."
            )
        source_file = _required_string(root, "source_file", "catalog")
        templates_raw = root.get("templates")
        partners_raw = root.get("partners")
        if not isinstance(templates_raw, list) or not templates_raw:
            raise SourceDataError(
                "Report template catalog must contain a non-empty templates list."
            )
        if not isinstance(partners_raw, list) or not partners_raw:
            raise SourceDataError("Report template catalog must contain a non-empty partners list.")

        catalog_directory = self.source_path.parent.resolve()
        templates: list[_CatalogTemplate] = []
        for index, value in enumerate(templates_raw, start=1):
            item = _mapping(value, f"template {index}")
            template_id = _required_string(item, "id", f"template {index}")
            name = _required_string(item, "name", f"template {index}")
            filename = _required_string(item, "file", f"template {index}")
            digest = _required_string(item, "sha256", f"template {index}").casefold()
            source_order = item.get("source_order")
            if not isinstance(source_order, int) or isinstance(source_order, bool):
                raise SourceDataError(f"Template '{template_id}' has an invalid source_order.")
            if not re.fullmatch(r"[a-f0-9]{64}", digest):
                raise SourceDataError(f"Template '{template_id}' has an invalid SHA-256 value.")
            relative_path = Path(filename)
            if relative_path.name != filename or relative_path.suffix.casefold() != ".docx":
                raise SourceDataError(
                    f"Template '{template_id}' must reference a DOCX file in the catalog directory."
                )
            template_path = (catalog_directory / relative_path).resolve()
            if template_path.parent != catalog_directory:
                raise SourceDataError(
                    f"Template '{template_id}' resolves outside the catalog directory."
                )
            templates.append(
                _CatalogTemplate(
                    id=template_id,
                    name=name,
                    source_order=source_order,
                    path=template_path,
                    sha256=digest,
                )
            )

        templates.sort(key=lambda template: template.source_order)
        _validate_unique_catalog_values(templates)
        expected_orders = list(range(1, len(templates) + 1))
        if [template.source_order for template in templates] != expected_orders:
            raise SourceDataError(
                "Template source_order values must be consecutive and start at 1."
            )

        partners: list[PartnerReference] = []
        for index, value in enumerate(partners_raw, start=1):
            item = _mapping(value, f"partner {index}")
            partners.append(
                PartnerReference(
                    name=_required_string(item, "name", f"partner {index}"),
                    icai_membership_number=_required_string(
                        item, "icai_membership_number", f"partner {index}"
                    ),
                )
            )
        if len({partner.name.casefold() for partner in partners}) != len(partners):
            raise SourceDataError("The report template catalog contains duplicate partner names.")
        if any(
            not re.fullmatch(r"\d{4,6}", partner.icai_membership_number) for partner in partners
        ):
            raise SourceDataError("The report template catalog contains an invalid ICAI number.")

        return _TemplateCatalog(source_file, tuple(templates), tuple(partners))

    def _read_document(self, source_path: Path, expected_sha256: str) -> DocumentObject:
        """Read a Word package reliably from OneDrive-backed and shared paths."""

        last_error: Exception | None = None
        for attempt in range(SOURCE_READ_ATTEMPTS):
            try:
                payload = source_path.read_bytes()
                actual_sha256 = hashlib.sha256(payload).hexdigest()
                if actual_sha256 != expected_sha256:
                    raise SourceDataError(
                        f"Standalone template '{source_path.name}' does not match its catalog "
                        "SHA-256. Regenerate the report template catalog."
                    )
                return Document(BytesIO(payload))
            except SourceDataError:
                raise
            except Exception as error:
                last_error = error
                if attempt + 1 < SOURCE_READ_ATTEMPTS:
                    time.sleep(SOURCE_READ_RETRY_SECONDS * (attempt + 1))
        if last_error is None:
            raise SourceDataError(f"Unable to read standalone template '{source_path.name}'.")
        raise SourceDataError(
            f"Unable to read standalone template '{source_path.name}': {last_error}"
        ) from last_error

    def _signature(self, catalog: _TemplateCatalog) -> tuple[_FileSignature, ...]:
        signatures = [self._file_signature(self.source_path, "Report template catalog")]
        signatures.extend(
            self._file_signature(template.path, "Standalone report template")
            for template in catalog.templates
        )
        return tuple(signatures)

    @staticmethod
    def _file_signature(path: Path, label: str) -> _FileSignature:
        try:
            stat = path.stat()
        except FileNotFoundError as error:
            raise SourceDataError(f"{label} is missing: {path}") from error
        except OSError as error:
            raise SourceDataError(
                f"Unable to inspect {label.lower()} '{path.name}': {error}"
            ) from error
        return _FileSignature(str(path), stat.st_mtime_ns, stat.st_size)

    @staticmethod
    def _extract_report_names(document: DocumentObject) -> tuple[tuple[str, str], ...]:
        """Extract ordered IDs and UI names from a master document's version matrix."""

        rows: list[tuple[str, str]] = []
        for table in document.tables:
            for row in table.rows:
                cells = [_clean(cell.text) for cell in row.cells]
                if len(cells) < 2 or not re.fullmatch(r"Report\s+[IVX]+", cells[0], re.I):
                    continue
                rows.append((_slug(cells[0]), cells[1]))
        return tuple(rows)

    @staticmethod
    def _extract_partners(document: DocumentObject) -> tuple[PartnerReference, ...]:
        """Extract partner names and ICAI numbers from a master document."""

        partners: list[PartnerReference] = []
        for table in document.tables:
            header_index = next(
                (
                    index
                    for index, row in enumerate(table.rows)
                    if any("icai membership number" in _normalized(cell.text) for cell in row.cells)
                ),
                None,
            )
            if header_index is None:
                continue
            for row in table.rows[header_index + 1 :]:
                cells = [_clean(cell.text) for cell in row.cells]
                if len(cells) < 2 or (cells[0] and re.fullmatch(r"\d+", cells[0])):
                    break
                name, membership_number = cells[-2:]
                if name and re.fullmatch(r"\d{4,6}", membership_number):
                    partners.append(PartnerReference(name, membership_number))
            if partners:
                break

        if len({partner.name.casefold() for partner in partners}) != len(partners):
            raise SourceDataError("The Word partner reference table contains duplicate names.")
        return tuple(partners)

    @staticmethod
    def _split_report_blocks(
        document: DocumentObject,
    ) -> tuple[tuple[Paragraph | Table, ...], ...]:
        groups: list[list[Paragraph | Table]] = []
        current: list[Paragraph | Table] | None = None
        for block in _iter_body_blocks(document):
            if isinstance(block, Paragraph) and _normalized(block.text) == REPORT_START:
                current = []
                groups.append(current)
            if current is not None:
                current.append(block)
        return tuple(tuple(group) for group in groups)

    @staticmethod
    def _extract_sections(
        report_id: str, blocks: tuple[Paragraph | Table, ...]
    ) -> tuple[ReportSection, ...]:
        sections: list[ReportSection] = []
        heading = "Report"
        heading_field_name: str | None = None
        paragraphs: list[ReportParagraph] = []
        heading_counts: dict[str, int] = {}

        def flush() -> None:
            nonlocal paragraphs
            body = "\n\n".join(
                _report_paragraph_text(paragraph)
                for paragraph in paragraphs
                if _paragraph_visible_by_default(paragraph)
            ).strip()
            if not body and sections:
                paragraphs = []
                return
            base = _slug(heading) or "section"
            occurrence = heading_counts.get(base, 0) + 1
            heading_counts[base] = occurrence
            suffix = f"-{occurrence}" if occurrence > 1 else ""
            sections.append(
                ReportSection(
                    id=f"{report_id}-{base}{suffix}",
                    heading=heading,
                    content=body,
                    order=len(sections) + 1,
                    paragraphs=tuple(paragraphs),
                    heading_field_name=heading_field_name,
                )
            )
            paragraphs = []

        for block_index, block in enumerate(blocks, start=1):
            if isinstance(block, Table):
                table_text = _table_text(block)
                if table_text:
                    paragraphs.append(
                        ReportParagraph(
                            id=f"{report_id}-p{block_index:03d}",
                            segments=(ReportContentSegment(table_text),),
                            order=block_index,
                            table=_extract_report_table(report_id, block_index, block),
                        )
                    )
                continue
            text = _clean(block.text)
            if not text:
                continue
            if _looks_like_heading(block, text):
                if paragraphs or sections:
                    flush()
                heading = text
                heading_field_name = (
                    _editable_field_name(report_id, block_index, 1)
                    if _paragraph_has_role(block, "editable")
                    else None
                )
            else:
                paragraphs.append(_extract_report_paragraph(report_id, block_index, block))
        if paragraphs or not sections:
            flush()
        return tuple(sections)


def _mapping(value: object, label: str) -> Mapping[str, object]:
    if not isinstance(value, dict) or not all(isinstance(key, str) for key in value):
        raise SourceDataError(f"Report template {label} must be a JSON object.")
    return value


def _required_string(item: Mapping[str, object], key: str, label: str) -> str:
    value = item.get(key)
    if not isinstance(value, str) or not value.strip():
        raise SourceDataError(f"Report template {label} must contain a non-empty '{key}'.")
    return value.strip()


def _validate_unique_catalog_values(templates: list[_CatalogTemplate]) -> None:
    checks = {
        "IDs": [template.id.casefold() for template in templates],
        "names": [template.name.casefold() for template in templates],
        "files": [template.path.name.casefold() for template in templates],
    }
    for label, values in checks.items():
        if len(set(values)) != len(values):
            raise SourceDataError(f"Report template catalog contains duplicate {label}.")


def _iter_body_blocks(document: DocumentObject) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _looks_like_heading(paragraph: Paragraph, text: str) -> bool:
    meaningful_runs = [run for run in paragraph.runs if run.text.strip()]
    if not meaningful_runs or len(text) > 180 or text.startswith("["):
        return False
    if text[-1:] in {".", ":", ";", ")"}:
        return False
    if not all(run.bold is True for run in meaningful_runs):
        return False
    excluded = (
        "for kkc",
        "chartered accountants",
        "firm registration",
        "icai membership",
        "udin",
        "place",
        "date",
    )
    return not _normalized(text).startswith(excluded)


def _table_text(table: Table) -> str:
    lines: list[str] = []
    for row in table.rows:
        values = [_clean(cell.text) for cell in row.cells]
        if any(values):
            lines.append(" | ".join(values))
    return "\n".join(lines)


def _extract_report_table(
    report_id: str,
    block_index: int,
    table: Table,
) -> ReportTable:
    cells: list[ReportTableCell] = []
    row_has_content: list[bool] = []
    column_count = max((len(row.cells) for row in table.rows), default=0)
    for row_index, row in enumerate(table.rows, start=1):
        has_content = False
        for column_index, cell in enumerate(row.cells, start=1):
            segments = _extract_table_cell_segments(
                report_id,
                block_index,
                row_index,
                column_index,
                cell.paragraphs,
            )
            has_content = has_content or any(segment.text.strip() for segment in segments)
            cells.append(
                ReportTableCell(
                    row=row_index,
                    column=column_index,
                    segments=segments,
                )
            )
        row_has_content.append(has_content)

    has_empty_body_rows = bool(row_has_content[:1]) and any(
        not has_content for has_content in row_has_content[1:]
    )
    repeatable_field_name = (
        f"word_{report_id.replace('-', '_')}_t{block_index:03d}_rows"
        if has_empty_body_rows and column_count > 0
        else None
    )
    return ReportTable(
        id=f"{report_id}-t{block_index:03d}",
        cells=tuple(cells),
        row_count=len(table.rows),
        column_count=column_count,
        repeatable_field_name=repeatable_field_name,
    )


def _extract_table_cell_segments(
    report_id: str,
    block_index: int,
    row_index: int,
    column_index: int,
    paragraphs: list[Paragraph],
) -> tuple[ReportContentSegment, ...]:
    merged: list[tuple[str, str]] = []
    for paragraph_index, paragraph in enumerate(paragraphs, start=1):
        if paragraph_index > 1 and merged and paragraph.text.strip():
            merged.append(("fixed", "\n"))
        for run in paragraph.runs:
            text = run.text.replace("\r", "").replace("\t", " ")
            if not text:
                continue
            role = _run_role(run)
            if merged and merged[-1][0] == role:
                previous_role, previous_text = merged[-1]
                merged[-1] = previous_role, previous_text + text
            else:
                merged.append((role, text))

    merged = _include_brackets_around_editable_text(merged)

    editable_index = 0
    segments: list[ReportContentSegment] = []
    for role, text in merged:
        field_name = None
        if role == "editable":
            editable_index += 1
            field_name = _known_editable_field(text) or (
                f"word_{report_id.replace('-', '_')}_t{block_index:03d}_"
                f"r{row_index:02d}_c{column_index:02d}_e{editable_index:02d}"
            )
        segments.append(ReportContentSegment(text=text, role=role, field_name=field_name))
    return tuple(segments)


def _extract_report_paragraph(
    report_id: str,
    block_index: int,
    paragraph: Paragraph,
) -> ReportParagraph:
    merged: list[tuple[str, str]] = []
    for run in paragraph.runs:
        text = run.text.replace("\r", "").replace("\t", " ")
        if not text:
            continue
        role = _run_role(run)
        if merged and merged[-1][0] == role:
            previous_role, previous_text = merged[-1]
            merged[-1] = previous_role, previous_text + text
        else:
            merged.append((role, text))

    merged = _include_brackets_around_editable_text(merged)

    while merged and not merged[0][1].strip():
        merged.pop(0)
    while merged and not merged[-1][1].strip():
        merged.pop()
    if merged:
        merged[0] = merged[0][0], merged[0][1].lstrip()
        merged[-1] = merged[-1][0], merged[-1][1].rstrip()

    editable_index = 0
    segments: list[ReportContentSegment] = []
    for role, text in merged:
        field_name = None
        if role == "editable":
            editable_index += 1
            field_name = _known_editable_field(text) or _editable_field_name(
                report_id, block_index, editable_index
            )
        segments.append(ReportContentSegment(text=text, role=role, field_name=field_name))

    instruction = _clean(
        "".join(segment.text for segment in segments if segment.role == "guidance")
    )
    visible_text = _clean(
        "".join(segment.text for segment in segments if segment.role != "guidance")
    )
    condition_field = None
    condition_value = None
    condition_default = None
    normalized_instruction = _normalized(instruction)
    if "removed in case data not received" in normalized_instruction:
        condition_field = f"word_data_received_{report_id.replace('-', '_')}"
        condition_value = "Yes"
        condition_default = "Yes"
    elif (
        "added in case data is not received" in normalized_instruction
        or "added in case data not received" in normalized_instruction
    ):
        condition_field = f"word_data_received_{report_id.replace('-', '_')}"
        condition_value = "No"
        condition_default = "Yes"
    elif instruction and visible_text:
        condition_field = f"word_include_{report_id.replace('-', '_')}_p{block_index:03d}"
        condition_value = "Yes"
        condition_default = "Yes"

    return ReportParagraph(
        id=f"{report_id}-p{block_index:03d}",
        segments=tuple(segments),
        order=block_index,
        condition_field=condition_field,
        condition_value=condition_value,
        condition_default=condition_default,
        instruction=instruction or None,
        numbering=_extract_paragraph_numbering(paragraph),
    )


def _include_brackets_around_editable_text(
    segments: list[tuple[str, str]],
) -> list[tuple[str, str]]:
    """Treat black brackets immediately surrounding blue text as editable text."""

    merged = list(segments)
    for index in range(1, len(merged) - 1):
        left_role, left_text = merged[index - 1]
        role, text = merged[index]
        right_role, right_text = merged[index + 1]
        if left_role != "fixed" or role != "editable" or right_role != "fixed":
            continue

        opening = re.search(r"(\[\s*)$", left_text)
        closing = re.match(r"(\s*\])", right_text)
        if opening is None or closing is None:
            continue

        merged[index - 1] = left_role, left_text[: opening.start()]
        merged[index] = role, opening.group(1) + text + closing.group(1)
        merged[index + 1] = right_role, right_text[closing.end() :]

    return [(role, text) for role, text in merged if text]


def _extract_paragraph_numbering(paragraph: Paragraph) -> ReportNumbering | None:
    paragraph_properties = paragraph._p.pPr
    numbering_properties = cast(
        Any, paragraph_properties.numPr if paragraph_properties is not None else None
    )
    if (
        numbering_properties is None
        or numbering_properties.numId is None
        or int(numbering_properties.numId.val) == 0
    ):
        return None

    numbering_id = int(numbering_properties.numId.val)
    level_index = int(numbering_properties.ilvl.val if numbering_properties.ilvl is not None else 0)
    numbering_root = cast(Any, paragraph.part).numbering_part.element
    numbering = next(
        (
            node
            for node in numbering_root.findall(qn("w:num"))
            if int(node.get(qn("w:numId"), "-1")) == numbering_id
        ),
        None,
    )
    if numbering is None:
        return None
    abstract_reference = numbering.find(qn("w:abstractNumId"))
    if abstract_reference is None:
        return None
    abstract_id = int(abstract_reference.get(qn("w:val"), "-1"))
    abstract = next(
        (
            node
            for node in numbering_root.findall(qn("w:abstractNum"))
            if int(node.get(qn("w:abstractNumId"), "-1")) == abstract_id
        ),
        None,
    )
    if abstract is None:
        return None

    start_overrides: dict[int, int] = {}
    for override in numbering.findall(qn("w:lvlOverride")):
        override_level = int(override.get(qn("w:ilvl"), "0"))
        start_override = override.find(qn("w:startOverride"))
        if start_override is not None:
            start_overrides[override_level] = int(start_override.get(qn("w:val"), "1"))

    levels: list[ReportNumberingLevel] = []
    level_elements = {
        int(node.get(qn("w:ilvl"), "0")): node for node in abstract.findall(qn("w:lvl"))
    }
    for current_level in range(level_index + 1):
        level = level_elements.get(current_level)
        if level is None:
            levels.append(ReportNumberingLevel(format="decimal", text="", start=1))
            continue
        format_element = level.find(qn("w:numFmt"))
        text_element = level.find(qn("w:lvlText"))
        start_element = level.find(qn("w:start"))
        start = start_overrides.get(
            current_level,
            int(start_element.get(qn("w:val"), "1")) if start_element is not None else 1,
        )
        levels.append(
            ReportNumberingLevel(
                format=(
                    format_element.get(qn("w:val"), "decimal")
                    if format_element is not None
                    else "decimal"
                ),
                text=(text_element.get(qn("w:val"), "") if text_element is not None else ""),
                start=start,
            )
        )

    direct_indent = (
        paragraph_properties.find(qn("w:ind")) if paragraph_properties is not None else None
    )
    level_properties = level_elements.get(level_index)
    level_indent = (
        level_properties.find(f"{qn('w:pPr')}/{qn('w:ind')}")
        if level_properties is not None
        else None
    )
    left_indent = _inherited_indent_value(direct_indent, level_indent, "left", "start")
    hanging_indent = _inherited_indent_value(direct_indent, level_indent, "hanging")
    if (
        direct_indent is not None
        and direct_indent.get(qn("w:firstLine")) is not None
        and direct_indent.get(qn("w:hanging")) is None
    ):
        hanging_indent = 0
    selected_level = level_elements.get(level_index)
    justification_element = (
        selected_level.find(qn("w:lvlJc")) if selected_level is not None else None
    )
    suffix_element = selected_level.find(qn("w:suff")) if selected_level is not None else None
    return ReportNumbering(
        id=numbering_id,
        level=level_index,
        levels=tuple(levels),
        left_indent_twips=left_indent,
        hanging_indent_twips=hanging_indent,
        justification=(
            justification_element.get(qn("w:val"), "left")
            if justification_element is not None
            else "left"
        ),
        suffix=(suffix_element.get(qn("w:val"), "tab") if suffix_element is not None else "tab"),
    )


def _inherited_indent_value(
    direct_indent: Any | None,
    level_indent: Any | None,
    *attribute_names: str,
) -> int:
    for indent in (direct_indent, level_indent):
        if indent is None:
            continue
        for attribute_name in attribute_names:
            value = indent.get(qn(f"w:{attribute_name}"))
            if value is not None:
                return int(value)
    return 0


def _run_role(run: object) -> str:
    font = getattr(run, "font", None)
    color = getattr(font, "color", None)
    rgb_value = getattr(color, "rgb", None)
    rgb = str(rgb_value).upper() if rgb_value is not None else ""
    if rgb == EDITABLE_BLUE_RGB:
        return "editable"
    if rgb in CONDITIONAL_GREEN_RGB:
        return "guidance"
    return "fixed"


def _paragraph_has_role(paragraph: Paragraph, role: str) -> bool:
    return any(run.text.strip() and _run_role(run) == role for run in paragraph.runs)


def _known_editable_field(text: str) -> str | None:
    normalized = re.sub(r"\s+", " ", text).strip().casefold()
    unwrapped = normalized.strip("[] ")
    if unwrapped in {"name of the company", "name of the bank"}:
        return "name_of_entity"
    if unwrapped == "name of the partner":
        return "partner_name"
    if re.fullmatch(r"x{6}", unwrapped):
        return "membership_number"
    if re.fullmatch(r"x{14}", unwrapped):
        return "udin"
    return None


def _editable_field_name(report_id: str, block_index: int, editable_index: int) -> str:
    return f"word_{report_id.replace('-', '_')}_p{block_index:03d}_e{editable_index:02d}"


def _report_paragraph_text(paragraph: ReportParagraph) -> str:
    return _clean(
        "".join(segment.text for segment in paragraph.segments if segment.role != "guidance")
    )


def _paragraph_visible_by_default(paragraph: ReportParagraph) -> bool:
    return (
        paragraph.condition_field is None
        or paragraph.condition_default == paragraph.condition_value
    )


def _clean(value: str) -> str:
    return re.sub(r"[ \t]+", " ", value.replace("\r", "")).strip()


def _normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold().replace("’", "'")


def _slug(value: str) -> str:
    normalized = _normalized(value)
    return re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
