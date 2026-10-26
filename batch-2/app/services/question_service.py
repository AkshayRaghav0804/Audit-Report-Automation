"""Load and normalize report questions from the source workbook."""

from __future__ import annotations

import ctypes
import logging
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.models.report_question import ReportQuestion, UnmappedQuestion
from app.models.report_template import ReportTemplate
from app.services.source_errors import SourceDataError

LOGGER = logging.getLogger(__name__)


@dataclass(frozen=True, slots=True)
class QuestionLoadResult:
    questions: tuple[ReportQuestion, ...]
    unmapped_rows: tuple[UnmappedQuestion, ...]
    common_question_count: int


@dataclass(frozen=True, slots=True)
class _FileSignature:
    modified_ns: int
    size: int


class QuestionService:
    """Parse the workbook, using common-question fallback when no report column exists."""

    def __init__(self, source_path: Path) -> None:
        self.source_path = source_path
        self._cached_key: tuple[_FileSignature, tuple[tuple[str, str], ...]] | None = None
        self._cached_result: QuestionLoadResult | None = None

    def load_questions(self, templates: tuple[ReportTemplate, ...]) -> QuestionLoadResult:
        signature = self._signature()
        template_key = tuple((template.id, template.name) for template in templates)
        key = signature, template_key
        if key == self._cached_key and self._cached_result is not None:
            return self._cached_result

        snapshot: tempfile.TemporaryDirectory[str] | None = None
        try:
            workbook, snapshot = _load_workbook_with_windows_snapshot(self.source_path)
        except Exception as error:
            raise SourceDataError(
                f"Unable to read question source '{self.source_path.name}': {error}"
            ) from error

        questions: list[ReportQuestion] = []
        unmapped: list[UnmappedQuestion] = []
        used_ids: set[str] = set()
        try:
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                header_index = self._find_header_row(rows)
                headers = {
                    _header_name(value): index
                    for index, value in enumerate(rows[header_index])
                    if value is not None
                }
                question_column = _first_column(headers, "question", "label", "question text")
                type_column = _first_column(headers, "type of answer", "field type", "type")
                if question_column is None:
                    continue
                report_column = _first_column(
                    headers, "report id", "template id", "report type", "report", "template"
                )
                description_column = _first_column(headers, "description", "help text")
                required_column = _first_column(headers, "required", "mandatory")
                options_column = _first_column(headers, "options", "choices")
                default_column = _first_column(headers, "default value", "default")
                section_column = _first_column(headers, "section id", "section")

                for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                    label = _cell_text(_at(row, question_column))
                    if not label:
                        continue
                    report_reference = _cell_text(_at(row, report_column))
                    report_id = self._map_report(report_reference, templates)
                    if report_reference and report_id is None:
                        unmapped.append(
                            UnmappedQuestion(
                                source_sheet=sheet.title,
                                source_row=row_number,
                                label=label,
                                report_reference=report_reference,
                                reason="No Word report ID or title matched this reference.",
                            )
                        )
                        continue

                    field_name = _slug(label)
                    question_id = _unique_id(field_name, used_ids)
                    description = _cell_text(_at(row, description_column))
                    raw_type = _cell_text(_at(row, type_column))
                    options_text = _cell_text(_at(row, options_column))
                    field_type = _field_type(raw_type)
                    questions.append(
                        ReportQuestion(
                            id=question_id,
                            report_id=report_id,
                            field_name=field_name.replace("-", "_"),
                            label=label,
                            field_type=field_type,
                            required=_boolean(_at(row, required_column)),
                            options=_options(field_type, options_text or description),
                            default_value=_cell_text(_at(row, default_column)) or None,
                            help_text=description or None,
                            order=row_number - header_index - 1,
                            section_id=_cell_text(_at(row, section_column)) or None,
                            source_sheet=sheet.title,
                            source_row=row_number,
                        )
                    )
        finally:
            workbook.close()
            if snapshot is not None:
                snapshot.cleanup()

        if not questions:
            raise SourceDataError(f"No usable questions were found in '{self.source_path.name}'.")
        result = QuestionLoadResult(
            questions=tuple(questions),
            unmapped_rows=tuple(unmapped),
            common_question_count=sum(question.report_id is None for question in questions),
        )
        self._cached_key = key
        self._cached_result = result
        LOGGER.info(
            "Loaded %d questions (%d common) from %s; %d rows unmapped",
            len(questions),
            result.common_question_count,
            self.source_path.name,
            len(unmapped),
        )
        return result

    def _signature(self) -> _FileSignature:
        try:
            stat = self.source_path.stat()
        except FileNotFoundError as error:
            raise SourceDataError(f"Question source file is missing: {self.source_path}") from error
        return _FileSignature(stat.st_mtime_ns, stat.st_size)

    @staticmethod
    def _find_header_row(rows: list[tuple[Any, ...]]) -> int:
        for index, row in enumerate(rows[:20]):
            headers = {_header_name(value) for value in row if value is not None}
            if "question" in headers:
                return index
        raise SourceDataError("The Excel source has no header row containing 'Question'.")

    @staticmethod
    def _map_report(report_reference: str, templates: tuple[ReportTemplate, ...]) -> str | None:
        if not report_reference:
            return None
        reference = _normalize(report_reference)
        for template in templates:
            if reference in {_normalize(template.id), _normalize(template.name)}:
                return template.id
        return None


def _at(row: tuple[Any, ...], column: int | None) -> Any:
    if column is None or column >= len(row):
        return None
    return row[column]


def _load_workbook_with_windows_snapshot(
    source_path: Path,
) -> tuple[Any, tempfile.TemporaryDirectory[str] | None]:
    """Open an Office/OneDrive workbook even when direct Python reads are denied."""

    try:
        return load_workbook(source_path, read_only=True, data_only=True), None
    except PermissionError as direct_error:
        if os.name != "nt":
            raise
        snapshot = tempfile.TemporaryDirectory(prefix="qc-question-source-")
        snapshot_path = Path(snapshot.name) / source_path.name
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        copied = kernel32.CopyFileW(str(source_path.resolve()), str(snapshot_path), False)
        if not copied:
            error_code = ctypes.get_last_error()
            snapshot.cleanup()
            raise PermissionError(
                f"Windows could not snapshot the open workbook (error {error_code})."
            ) from direct_error
        try:
            workbook = load_workbook(snapshot_path, read_only=True, data_only=True)
        except Exception:
            snapshot.cleanup()
            raise
        return workbook, snapshot


def _first_column(headers: dict[str, int], *names: str) -> int | None:
    return next((headers[name] for name in names if name in headers), None)


def _header_name(value: Any) -> str:
    return _normalize(_cell_text(value))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-") or "question"


def _unique_id(base: str, used: set[str]) -> str:
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}-{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _field_type(value: str) -> str:
    normalized = _normalize(value)
    return {
        "text": "text",
        "number": "number",
        "date": "date",
        "yes no": "yes_no",
        "boolean": "yes_no",
        "multiple choice": "select",
        "select": "select",
        "multiselect": "multiselect",
        "multi select": "multiselect",
        "checkbox": "checkbox",
        "radio": "radio",
        "long text": "textarea",
        "textarea": "textarea",
    }.get(normalized, "text")


def _boolean(value: Any) -> bool:
    return _normalize(_cell_text(value)) in {"yes", "true", "required", "1"}


def _options(field_type: str, value: str) -> tuple[str, ...]:
    if field_type == "yes_no":
        return "Yes", "No"
    if field_type != "select":
        return ()
    selected = re.search(r"\bselect\s*:\s*(.+)", value, flags=re.I)
    choices = selected.group(1) if selected else value
    return tuple(
        choice.strip().rstrip(".") for choice in re.split(r"[,;]", choices) if choice.strip()
    )
