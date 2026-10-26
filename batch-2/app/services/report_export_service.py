"""Generate stage-aware Word and PDF snapshots from controlled report sources."""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models.report_template import (
    ReportContentSegment,
    ReportSection,
    ReportTable,
    ReportTemplate,
)
from app.services.report_basic_details import report_text_for_answers
from app.services.report_draft_service import (
    SOURCE_PLACEHOLDERS,
    format_answer_for_preview,
    link_section_override_fields,
    section_placeholder_texts,
    source_paragraph_is_visible,
    source_section_is_visible,
    substitute_placeholders,
)
from app.source_config import ReportSourceConfig

EDITABLE_BLUE_RGB = "00B0F0"
CONDITIONAL_GREEN_RGB = frozenset({"4EA72E", "92D050"})
BLACK_RGB = RGBColor(0, 0, 0)
EDITABLE_BLUE_COLOR = RGBColor(0, 176, 240)


class ReportExportError(RuntimeError):
    """Raised when a report snapshot cannot be generated."""


@dataclass(frozen=True, slots=True)
class ReportExportArtifact:
    content: bytes
    filename: str
    media_type: str


class ReportExportService:
    def __init__(self, template_directory: Path | None = None) -> None:
        config = ReportSourceConfig.from_environment()
        self.template_directory = template_directory or config.template_path.parent

    def export_docx(
        self,
        template: ReportTemplate,
        answers: dict[str, Any],
        section_overrides: dict[str, str],
        *,
        heading_overrides: dict[str, str] | None = None,
        stage: str,
        report_id: str | None,
        highlighted: bool = False,
    ) -> ReportExportArtifact:
        document = self._build_document(
            template,
            answers,
            section_overrides,
            heading_overrides=heading_overrides,
            stage=stage,
            highlighted=highlighted,
        )
        stream = BytesIO()
        document.save(stream)
        return ReportExportArtifact(
            content=stream.getvalue(),
            filename=_export_filename(
                template, answers, stage, report_id, "docx", highlighted=highlighted
            ),
            media_type=("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        )

    def export_pdf(
        self,
        template: ReportTemplate,
        answers: dict[str, Any],
        section_overrides: dict[str, str],
        *,
        heading_overrides: dict[str, str] | None = None,
        stage: str,
        report_id: str | None,
        highlighted: bool = False,
    ) -> ReportExportArtifact:
        word_artifact = self.export_docx(
            template,
            answers,
            section_overrides,
            heading_overrides=heading_overrides,
            stage=stage,
            report_id=report_id,
            highlighted=highlighted,
        )
        # Microsoft Word can be blocked from automating files in Local Temp or hidden
        # directories by enterprise Office policies. The project's normal output location
        # is trusted on the target workstation and each conversion directory is short-lived.
        export_work_root = self.template_directory.parent / "outputs"
        export_work_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
            prefix="qc_report_export_",
            dir=export_work_root,
            ignore_cleanup_errors=True,
        ) as temporary_directory:
            workdir = Path(temporary_directory)
            docx_path = workdir / "report.docx"
            pdf_path = workdir / "report.pdf"
            docx_path.write_bytes(word_artifact.content)
            _convert_docx_to_pdf(docx_path, pdf_path)
            try:
                pdf_content = pdf_path.read_bytes()
            except FileNotFoundError as error:
                raise ReportExportError("PDF conversion did not produce an output file.") from error
        if not pdf_content.startswith(b"%PDF"):
            raise ReportExportError("The generated PDF file is invalid.")
        return ReportExportArtifact(
            content=pdf_content,
            filename=_export_filename(
                template, answers, stage, report_id, "pdf", highlighted=highlighted
            ),
            media_type="application/pdf",
        )

    def _build_document(
        self,
        template: ReportTemplate,
        answers: dict[str, Any],
        section_overrides: dict[str, str],
        *,
        heading_overrides: dict[str, str] | None,
        stage: str,
        highlighted: bool,
    ) -> DocumentObject:
        source_path = self.template_directory / template.source_file
        try:
            document = Document(BytesIO(source_path.read_bytes()))
        except (OSError, ValueError) as error:
            raise ReportExportError(
                f"Unable to open the controlled Word source '{template.source_file}'."
            ) from error

        document.core_properties.title = template.name
        document.core_properties.subject = f"QC Report Automation - {stage}"
        _remove_print_only_instructions(document)
        blocks = list(_iter_body_blocks(document))
        heading_indexes = _section_heading_indexes(blocks, template.sections)
        paragraphs_by_order = {
            paragraph.order: paragraph
            for section in template.sections
            for paragraph in section.paragraphs
        }

        for section_index, section in reversed(tuple(enumerate(template.sections))):
            heading_index = heading_indexes[section_index]
            end_index = (
                heading_indexes[section_index + 1]
                if section_index + 1 < len(heading_indexes)
                else len(blocks)
            )
            section_blocks = blocks[heading_index:end_index]
            if not source_section_is_visible(section, answers):
                for block in section_blocks:
                    _remove_block(block)
                continue

            heading_block = blocks[heading_index]
            if isinstance(heading_block, Paragraph):
                heading_override = (heading_overrides or {}).get(section.id)
                if heading_override is not None:
                    _replace_paragraph_text(heading_block, heading_override)
                else:
                    heading_segments = (
                        (
                            ReportContentSegment(
                                section.heading,
                                role="editable",
                                field_name=section.heading_field_name,
                            ),
                        )
                        if section.heading_field_name
                        else ()
                    )
                    _patch_paragraph(
                        heading_block,
                        heading_segments,
                        answers,
                        highlighted=highlighted,
                    )

            override = section_overrides.get(section.id)
            if override is not None:
                for block in section_blocks[1:]:
                    _remove_block(block)
                if isinstance(heading_block, Paragraph):
                    linked_override = link_section_override_fields(section, override)
                    _insert_override_paragraph(
                        heading_block,
                        substitute_placeholders(
                            linked_override,
                            answers,
                            placeholder_texts=section_placeholder_texts(section),
                        ),
                    )
                continue

            for relative_index, block in enumerate(section_blocks[1:], start=1):
                absolute_index = heading_index + relative_index
                block_index = absolute_index - heading_indexes[0] + 1
                paragraph = paragraphs_by_order.get(block_index)
                if paragraph is None:
                    continue
                if not source_paragraph_is_visible(paragraph, answers):
                    _remove_block(block)
                    continue
                if isinstance(block, Paragraph):
                    _patch_paragraph(block, paragraph.segments, answers, highlighted=highlighted)
                elif isinstance(block, Table) and paragraph.table is not None:
                    _patch_table(block, paragraph.table, answers, highlighted=highlighted)
        _finalize_source_colors(document, highlighted=highlighted)
        return document


def _remove_print_only_instructions(document: DocumentObject) -> None:
    """Remove source-editor directions which must not appear in report snapshots."""
    paragraphs = list(document.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        if "delete before printing" not in paragraph.text.casefold():
            continue
        for run in paragraph.runs:
            run.text = ""


def _finalize_source_colors(document: DocumentObject, *, highlighted: bool) -> None:
    """Make snapshots print-ready even for source runs outside parsed report sections."""
    containers: list[Any] = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in _container_paragraphs(container):
            for run in paragraph.runs:
                role = _run_role(run)
                if role == "guidance":
                    run.text = ""
                elif role == "editable":
                    run.font.color.rgb = EDITABLE_BLUE_COLOR if highlighted else BLACK_RGB
    _finalize_xml_source_colors(document, highlighted=highlighted)


def _finalize_xml_source_colors(document: DocumentObject, *, highlighted: bool) -> None:
    """Cover Word runs nested in AlternateContent which python-docx does not expose."""
    roots = [document.element]
    for section in document.sections:
        roots.extend((section.header._element, section.footer._element))
    for root in roots:
        for run_element in root.iter(qn("w:r")):
            color_values = {
                str(color.get(qn("w:val"), "")).upper() for color in run_element.iter(qn("w:color"))
            }
            if color_values.intersection(CONDITIONAL_GREEN_RGB):
                for text_element in run_element.iter(qn("w:t")):
                    text_element.text = ""
        for color in root.iter(qn("w:color")):
            value = str(color.get(qn("w:val"), "")).upper()
            if value == EDITABLE_BLUE_RGB:
                color.set(qn("w:val"), EDITABLE_BLUE_RGB if highlighted else "000000")
            elif value in CONDITIONAL_GREEN_RGB:
                color.set(qn("w:val"), "000000")


def _container_paragraphs(container: Any) -> tuple[Paragraph, ...]:
    paragraphs = list(container.paragraphs)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(_container_paragraphs(cell))
    return tuple(paragraphs)


def _iter_body_blocks(document: DocumentObject) -> tuple[Paragraph | Table, ...]:
    blocks: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            blocks.append(Paragraph(child, document))
        elif isinstance(child, CT_Tbl):
            blocks.append(Table(child, document))
    return tuple(blocks)


def _section_heading_indexes(
    blocks: list[Paragraph | Table],
    sections: tuple[ReportSection, ...],
) -> tuple[int, ...]:
    indexes: list[int] = []
    cursor = 0
    for section in sections:
        normalized_heading = _normalize(section.heading)
        match = None
        for index in range(cursor, len(blocks)):
            block = blocks[index]
            if isinstance(block, Paragraph) and _normalize(block.text) == normalized_heading:
                match = index
                break
        if match is None:
            raise ReportExportError(
                f"Unable to locate the Word section '{section.heading}' for export."
            )
        indexes.append(match)
        cursor = match + 1
    return tuple(indexes)


def _patch_paragraph(
    paragraph: Paragraph,
    segments: tuple[ReportContentSegment, ...],
    answers: dict[str, Any],
    *,
    highlighted: bool,
) -> None:
    editable_segments = tuple(segment for segment in segments if segment.role == "editable")
    editable_groups = _run_groups(paragraph, "editable")
    for group, segment in zip(editable_groups, editable_segments, strict=False):
        value = (
            answers.get(segment.field_name, segment.text) if segment.field_name else segment.text
        )
        _replace_editable_run_group(paragraph, group, format_answer_for_preview(value))
        for run in group:
            run.font.color.rgb = EDITABLE_BLUE_COLOR if highlighted else BLACK_RGB

    for group in _run_groups(paragraph, "guidance"):
        _replace_run_group(group, "")
    _replace_fixed_placeholders(paragraph, answers)
    _replace_managed_basic_details(paragraph, answers)


def _replace_paragraph_text(paragraph: Paragraph, value: str) -> None:
    """Replace a paragraph while retaining its first run's controlled formatting."""

    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _run_groups(paragraph: Paragraph, target_role: str) -> tuple[tuple[Any, ...], ...]:
    groups: list[list[Any]] = []
    active: list[Any] = []
    for run in paragraph.runs:
        role = _run_role(run)
        if role == target_role:
            active.append(run)
        elif active:
            groups.append(active)
            active = []
    if active:
        groups.append(active)
    return tuple(tuple(group) for group in groups)


def _replace_run_group(group: tuple[Any, ...], value: str) -> None:
    if not group:
        return
    group[0].text = value
    for run in group[1:]:
        run.text = ""


def _replace_editable_run_group(
    paragraph: Paragraph,
    group: tuple[Any, ...],
    value: str,
) -> None:
    """Replace blue text and consume paired black brackets that belong to it."""

    runs = list(paragraph.runs)
    if group:
        group_element_ids = {id(run._r) for run in group}
        group_indexes = [index for index, run in enumerate(runs) if id(run._r) in group_element_ids]
        if group_indexes:
            left_run = _nearest_text_run(runs, group_indexes[0] - 1, step=-1)
            right_run = _nearest_text_run(runs, group_indexes[-1] + 1, step=1)
            if left_run is not None and right_run is not None:
                opening = re.search(r"(\[\s*)$", left_run.text)
                closing = re.match(r"(\s*\])", right_run.text)
                if opening is not None and closing is not None:
                    left_run.text = left_run.text[: opening.start()]
                    right_run.text = right_run.text[closing.end() :]
    _replace_run_group(group, value)


def _nearest_text_run(runs: list[Any], start: int, *, step: int) -> Any | None:
    index = start
    while 0 <= index < len(runs):
        if runs[index].text:
            return runs[index]
        index += step
    return None


def _replace_fixed_placeholders(paragraph: Paragraph, answers: dict[str, Any]) -> None:
    replacements = {
        placeholder: format_answer_for_preview(answers[field_name])
        for field_name, placeholders in SOURCE_PLACEHOLDERS.items()
        if field_name in answers and format_answer_for_preview(answers[field_name])
        for placeholder in placeholders
    }
    if not replacements:
        return
    for run in paragraph.runs:
        if _run_role(run) != "fixed":
            continue
        value = run.text
        for placeholder, replacement in replacements.items():
            value = value.replace(placeholder, replacement)
        run.text = value


def _replace_managed_basic_details(paragraph: Paragraph, answers: dict[str, Any]) -> None:
    for group in _run_groups(paragraph, "fixed"):
        source = "".join(run.text for run in group)
        rendered = report_text_for_answers(source, answers)
        if rendered != source:
            _replace_run_group(group, rendered)


def _patch_table(
    table: Table,
    source_table: ReportTable,
    answers: dict[str, Any],
    *,
    highlighted: bool,
) -> None:
    cells = {(cell.row, cell.column): cell for cell in source_table.cells}
    for row_index, row in enumerate(table.rows, start=1):
        for column_index, cell in enumerate(row.cells, start=1):
            source_cell = cells.get((row_index, column_index))
            if source_cell is None:
                continue
            editable_segments = tuple(
                segment for segment in source_cell.segments if segment.role == "editable"
            )
            editable_index = 0
            for paragraph in cell.paragraphs:
                groups = _run_groups(paragraph, "editable")
                for group in groups:
                    if editable_index >= len(editable_segments):
                        break
                    segment = editable_segments[editable_index]
                    value = (
                        answers.get(segment.field_name, segment.text)
                        if segment.field_name
                        else segment.text
                    )
                    _replace_editable_run_group(
                        paragraph,
                        group,
                        format_answer_for_preview(value),
                    )
                    for run in group:
                        run.font.color.rgb = EDITABLE_BLUE_COLOR if highlighted else BLACK_RGB
                    editable_index += 1
                for guidance_group in _run_groups(paragraph, "guidance"):
                    _replace_run_group(guidance_group, "")
                _replace_fixed_placeholders(paragraph, answers)
                _replace_managed_basic_details(paragraph, answers)

    if not source_table.repeatable_field_name or len(table.rows) < 2:
        return
    values = _normalize_rows(
        answers.get(source_table.repeatable_field_name),
        source_table.column_count,
    )
    template_row = deepcopy(table.rows[1]._tr)
    for row in tuple(table.rows[1:]):
        table._tbl.remove(row._tr)
    for row_values in values:
        table._tbl.append(deepcopy(template_row))
        new_row = table.rows[-1]
        for column_index, value in enumerate(row_values):
            if column_index < len(new_row.cells):
                new_row.cells[column_index].text = value
                if highlighted:
                    for paragraph in new_row.cells[column_index].paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = EDITABLE_BLUE_COLOR


def _normalize_rows(value: Any, column_count: int) -> tuple[tuple[str, ...], ...]:
    if not isinstance(value, list):
        return ()
    rows: list[tuple[str, ...]] = []
    for item in value:
        if not isinstance(item, (list, tuple)):
            continue
        row = [str(cell or "") for cell in item[:column_count]]
        row.extend([""] * (column_count - len(row)))
        if any(cell.strip() for cell in row):
            rows.append(tuple(row))
    return tuple(rows)


def _run_role(run: Any) -> str:
    rgb_value = getattr(getattr(run.font, "color", None), "rgb", None)
    rgb = str(rgb_value).upper() if rgb_value is not None else ""
    if rgb == EDITABLE_BLUE_RGB:
        return "editable"
    if rgb in CONDITIONAL_GREEN_RGB:
        return "guidance"
    return "fixed"


def _insert_override_paragraph(heading: Paragraph, content: str) -> None:
    paragraph_xml = OxmlElement("w:p")
    heading._p.addnext(paragraph_xml)
    paragraph = Paragraph(paragraph_xml, heading._parent)
    paragraph.add_run(content)


def _remove_block(block: Paragraph | Table) -> None:
    element = block._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold().replace("â€™", "'")


def _export_filename(
    template: ReportTemplate,
    answers: dict[str, Any],
    stage: str,
    report_id: str | None,
    extension: str,
    *,
    highlighted: bool = False,
) -> str:
    entity = format_answer_for_preview(answers.get("name_of_entity")) or template.id
    parts = (entity, report_id or "draft", stage, "Blue Highlights" if highlighted else None)
    stem = " - ".join(_safe_filename_part(part) for part in parts if part)
    return f"{stem}.{extension}"


def _safe_filename_part(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "-", str(value)).strip(" .-")
    return re.sub(r"\s+", " ", cleaned)[:80] or "report"


def _convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        completed = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pdf_path.parent),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
        )
        converted = pdf_path.parent / f"{docx_path.stem}.pdf"
        if completed.returncode == 0 and converted.exists():
            if converted != pdf_path:
                converted.replace(pdf_path)
            return

    powershell_script = """
$word = $null
$document = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $sourcePath = $env:QC_REPORT_DOCX_PATH
    $pdfPath = $env:QC_REPORT_PDF_PATH
    $pdfFormat = 17
    $document = $word.Documents.Open($sourcePath)
    $document.SaveAs2([ref]$pdfPath, [ref]$pdfFormat)
}
finally {
    if ($null -ne $document) { $document.Close($false) }
    if ($null -ne $word) { $word.Quit() }
}
"""
    try:
        completed = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-Command",
                powershell_script,
            ],
            capture_output=True,
            text=True,
            timeout=90,
            check=False,
            env={
                **os.environ,
                "QC_REPORT_DOCX_PATH": str(docx_path.resolve()),
                "QC_REPORT_PDF_PATH": str(pdf_path.resolve()),
            },
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ReportExportError(
            "PDF export requires Microsoft Word or LibreOffice on this computer."
        ) from error
    if completed.returncode != 0 or not pdf_path.exists():
        detail = (completed.stderr or completed.stdout).strip()
        raise ReportExportError(
            "PDF export requires Microsoft Word or LibreOffice on this computer."
            + (f" Conversion detail: {detail}" if detail else "")
        )
